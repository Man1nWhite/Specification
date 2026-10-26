"""Static packaging checks, headless VBA preview test, and Python generator tests."""

import re
import shutil
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import docx
import openpyxl
from docx.shared import Pt
from lxml import etree
from pyopenvba import ExcelFile
from pyopenvba.apps.excel import ExcelApplication

from generate_tkp import (
    _eval_cell,
    build_preview_rows,
    clean_text,
    find_default_template,
    format_number,
    format_qty,
    generate_tkp,
    read_tkp_data,
)
from tools import sync_assets

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "Gen_TKP.xlsm"
TEMPLATE = ROOT / "ТКП_шаблон.dotx"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": W}


def normalise_preview_rows(rows):
    """Normalise preview rows from either engine so they can be compared.

    Excel returns every number as a float, so an integral quantity comes back as
    2.0 where the Python generator builds 2. Collapsing integral floats keeps the
    comparison faithful, unlike an int() cast, which would silently turn 2.9 into
    2 and hide a real difference between the engines.
    """
    normalised = []
    for num, description, article, unit, qty, marker in rows:
        if isinstance(qty, float) and qty.is_integer():
            qty = int(qty)
        normalised.append((num, description, article or "", unit, qty, marker))
    return normalised


def read_part(name: str):
    """Parse one part of the shipped template."""
    with ZipFile(TEMPLATE) as archive:
        return etree.fromstring(archive.read(name))


def element_text(element) -> str:
    return "".join(t.text or "" for t in element.iter(f"{{{W}}}t"))


def template_body():
    body = read_part("word/document.xml").find(f"{{{W}}}body")
    assert body is not None
    return body


def rows_of(table) -> list[list[str]]:
    """The cell texts of every row; a merged row yields a single cell."""
    return [
        [
            " ".join("".join(p.itertext()) for p in cell.findall(f"{{{W}}}p")).strip()
            for cell in row.findall(f"{{{W}}}tc")
        ]
        for row in table.findall(f"{{{W}}}tr")
    ]


def find_table(body, needle: str):
    """The table under `body` whose text mentions `needle`."""
    for table in body.findall(f"{{{W}}}tbl"):
        if needle in element_text(table):
            return table
    raise AssertionError(f"No table contains {needle!r}")


class TKPTest(unittest.TestCase):
    def test_embedded_source_matches_editable_source(self):
        expected = (ROOT / "Module1.bas").read_text(encoding="utf-8").replace("\r\n", "\n").strip()
        with ExcelFile(WORKBOOK) as book:
            actual = book.get_module("Module1").replace("\r\n", "\n").strip()
            self.assertEqual(actual, expected)
            self.assertIn("Module12", book.module_names())  # existing macros preserved
        self.assertIn('"ТКП_шаблон.dotx"', expected)
        self.assertNotIn("ТКП_шаблон.docx", expected)
        self.assertNotIn("Коммерческое предложение_шаблон.docx", expected)
        self.assertNotIn("Коммерческое предложение.docx", expected)

    def test_bookmarks_are_in_their_destination_not_the_notes(self):
        self.assertTrue(TEMPLATE.exists())
        with ZipFile(TEMPLATE) as archive:
            document = etree.fromstring(archive.read("word/document.xml"))
            header = etree.fromstring(archive.read("word/header1.xml"))
        body = document.find(".//w:body", NS)
        self.assertIsNotNone(body)
        notes = next(
            p
            for p in body.findall("w:p", NS)
            if "Примечания к Приложению" in "".join(p.itertext())
        )
        target = {}
        for name in ("Num", "Customer", "Object", "Subject", "Spec_Table"):
            starts = document.xpath(
                ".//w:bookmarkStart[@w:name=$name]", namespaces=NS, name=name
            )
            self.assertEqual(len(starts), 1)
            self.assertIsNot(starts[0].getparent(), notes)
            target[name] = starts[0]
        self.assertLess(body.index(target["Spec_Table"].getparent()), body.index(notes))
        self.assertEqual(
            len(header.xpath(".//w:bookmarkStart[@w:name='HeaderNum']", namespaces=NS)),
            1,
        )
        for name in ("Num", "Customer", "Object", "Subject"):
            self.assertEqual(target[name].getparent().tag, f"{{{W}}}p")
            self.assertEqual(target[name].getparent().getparent().tag, f"{{{W}}}tc")

    def test_bookmark_ids_are_unique_and_paired(self):
        """Every bookmark opens once and closes once.

        Regression: the sync step created its bookmarks with hard-coded ids
        (24, 25, 26, 29) that the template already used for its `_Hlk` and `_Toc`
        marks, so one id stood for two bookmarks and Word offered to repair the
        file.
        """
        document = read_part("word/document.xml")
        starts = [start.get(f"{{{W}}}id") for start in document.iter(f"{{{W}}}bookmarkStart")]
        ends = [end.get(f"{{{W}}}id") for end in document.iter(f"{{{W}}}bookmarkEnd")]
        self.assertEqual(len(starts), len(set(starts)))
        self.assertEqual(len(ends), len(set(ends)))
        self.assertEqual(set(starts), set(ends))

    def test_sync_leaves_a_synchronised_template_untouched(self):
        """The sync step is idempotent on the template it ships.

        The reviewed edition of the template is the source of truth: its short
        blocks are the decision, not a loss to be restored. Running the step
        again must therefore change nothing — byte for byte, in every part.
        """
        with TemporaryDirectory() as tmp:
            copy = Path(tmp) / TEMPLATE.name
            shutil.copy(TEMPLATE, copy)
            original = sync_assets.TEMPLATE_DOTX
            sync_assets.TEMPLATE_DOTX = copy
            try:
                sync_assets.patch_template()
                with ZipFile(copy) as archive:
                    first = {name: archive.read(name) for name in archive.namelist()}
                sync_assets.patch_template()
                with ZipFile(copy) as archive:
                    second = {name: archive.read(name) for name in archive.namelist()}
            finally:
                sync_assets.TEMPLATE_DOTX = original
        self.assertEqual(sorted(first), sorted(second))
        self.assertEqual(
            [name for name in sorted(first) if first[name] != second[name]], []
        )

    def test_default_template_is_only_the_dotx_asset(self):
        self.assertEqual(find_default_template(ROOT), TEMPLATE)
        with TemporaryDirectory() as tmp:
            base_dir = Path(tmp)
            for name in (
                "ТКП_шаблон.docx",
                "Коммерческое предложение_шаблон.docx",
                "Коммерческое предложение.docx",
            ):
                (base_dir / name).touch()
            with self.assertRaises(FileNotFoundError):
                find_default_template(base_dir)

    def test_template_has_no_latin_letters_inside_russian_words(self):
        """A Latin "e" typed instead of "ё" is invisible to the spellchecker.

        The template used to ship "Объeм", "выдаeтся", "приeмо-сдаточные",
        "надeжность", "счeт" and "трeхуровневой" — every one of them reaches the
        customer as a typo because Word does not flag a Latin letter inside a
        Russian word. Only the certificate number "RU.AЖ49.H04589" mixes scripts.
        """
        cyrillic = re.compile(r"[Ѐ-ӿ]")
        latin = re.compile(r"[A-Za-z]")
        mixed = set()
        for name in ("word/document.xml", "word/header1.xml", "word/header2.xml"):
            with ZipFile(TEMPLATE) as archive:
                root = etree.fromstring(archive.read(name))
            text = " ".join(t.text or "" for t in root.iter(f"{{{W}}}t"))
            for token in re.findall(r"[^\W\d_]+", text):
                if cyrillic.search(token) and latin.search(token):
                    mixed.add((name, token))
        self.assertEqual(mixed, {("word/document.xml", "AЖ")})

    def test_template_uses_the_proofread_wording(self):
        """The wording follows the reference ТКП: «ё», «работ», «Акта», «тайну»."""
        with ZipFile(TEMPLATE) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        text = "".join(t.text or "" for t in root.iter(f"{{{W}}}t"))
        for broken in (
            "Оплата работы по пунктам",
            "стоимости работы по этапу",
            "Комплексного опробование",
            "третьих сторон и лиц",
            "оформлением Акт о завершении",
            "исходных данных, Спецификация",
            "Приемо-сдаточные испытания",
            "приема-передачи",
            "приемки-передачи",
            "3 (трех)",
            "собственности Заказчику",
            "третьим лицам или сторонам",
            "Монтаж КИПиА",
            "Монтаж и подключение кабельных линий",
            "Для управления следующими системами",
            "контроллеров шкафов диспетчеризации",
            "Исполнительная документация",
            "оперативного и инженеров",
            "заземляющими элементами",
            "распределенных",
            "включена в состав средств измерения Госреестра",
            "пуско-наладочных",
            "ГОСТ Р ИСО 9001 – 2015",
            "ГОСТ Р ИСО 45001 – 2020",
            "~220В",
            "производителей, номенклатура",
            "Абельный журнал",
            "21 неделя",
            "3 (три) дня",
            "4 (четыре) дня",
            "Срок начала работ согласовывается Сторонами в договоре",
            "Дата начала выполнения работ согласовывается в договоре",
            # This review: the vague note, the three-day advance, the works
            # inside the delivery terms, the unlinked МИРТС, the old heading and
            # the two rows that duplicated the specific documents.
            "могут быть уточнены",
            "в течение 3 (трёх) календарных дней",
            "и выполнение работ по созданию АСУ ТП осуществляется",
            "Система «МИРТС» является основой",
            "Допущения:",
            "Документация на ППО",
            "Эксплуатационная документация (руководство оператора, паспорта)",
        ):
            self.assertNotIn(broken, text)
        for fixed in (
            "Оплата работ по пунктам 5–8",
            "от стоимости работ по этапу",
            "Комплексное опробование",
            "для третьих лиц",
            "оформлением Акта о завершении ПНР",
            "При изменении исходных данных Спецификация",
            "Объём поставки",
            "выдаётся сертификат",
            "надёжность",
            "за счёт трёхуровневой",
            "собственности к Заказчику",
            "для создания АСУ ТП осуществляется на склад Заказчика по адресу:",
            "Технико-коммерческое предложение",
            "Перечень передаваемой документации",
            "Кабельный журнал",
            "Исполнительные схемы",
            "Акты скрытых работ",
            "Руководство оператора",
            "Руководство администратора",
            "Паспорта и сертификаты соответствия на покупное оборудование",
            "Акт выполненных пусконаладочных работ",
            "Настройка контроллеров ПТК",
            "оперативного персонала и инженеров АСУ ТП",
            "10 (десять) недель с момента окончания работ по этапу № 6",
            "Монтаж оборудования",
            "между заземлителями и заземляемыми элементами",
            "в электроустановках напряжением до 1000 В",
            "распределённых АСУ ТП",
            "внесена в Госреестр средств измерений",
            "пусконаладочных, интеграционных услуг",
            "ГОСТ Р ИСО 9001-2015",
            "ГОСТ Р ИСО 45001-2020",
            "~220 В, =220 В, =24 В",
            "третьим лицам без письменного разрешения",
            "Маркировка шкафов и оборудования",
            "Подключение питания шкафов",
            "Заземление шкафов",
            "22 недели",
            "7 (семь) дней с момента окончания работ по этапу № 3",
            "7 (семь) дней с момента окончания работ по этапу № 4",
            "Дата проведения обучения согласуется заранее",
            "Оплату командировочных расходов (трансфер до Москвы и обратно, "
            "проживание, питание) Заказчик осуществляет самостоятельно",
            # Payment terms: the second payment and the title pass with the УПД,
            # the stage advance is seven days, the invoice follows the act.
            "универсального передаточного документа (УПД)",
            "7 (семи) календарных дней с даты начала этапа",
            "Возможен зачёт аванса",
            "счёт выставляется после подписания Сторонами акта оказанных услуг",
            # The certification now ties АРКС400 to the measuring system.
            "ПТК АРКС400 входит в состав системы «МИРТС»",
            # The whole project has a duration, on page two and in the chart.
            "Общий срок реализации проекта",
            "35 недель",
        ):
            self.assertIn(fixed, text)

    def test_template_keeps_compound_words_on_one_line(self):
        """A hyphen at the end of a line vanishes in PDF text extraction.

        «Акта приёма-передачи» used to be read as «приёмапередачи» because the
        compound was split across two lines in a narrow table cell. A
        non-breaking hyphen keeps the word in one piece.
        """
        with ZipFile(TEMPLATE) as archive:
            root = etree.fromstring(archive.read("word/document.xml"))
        text = "".join(t.text or "" for t in root.iter(f"{{{W}}}t"))
        self.assertIn("Акта приёма‑передачи документации", text)
        self.assertIn("Протокола приёма‑передачи", text)
        self.assertIn("Приёмо‑сдаточные испытания", text)
        self.assertIn("приёмо‑сдаточных испытаний", text)
        self.assertIn("приёмо‑сдаточные испытания", text)
        # No plain hyphen is left inside these compounds.
        for plain in ("приёма-передачи", "приёмо-сдаточные", "приёмо-сдаточных"):
            self.assertNotIn(plain, text)

    def test_template_document_title_is_tkp(self):
        """Word names a new document after the template's title property."""
        with ZipFile(TEMPLATE) as archive:
            core = etree.fromstring(archive.read("docProps/core.xml"))
            app = etree.fromstring(archive.read("docProps/app.xml"))
        title = core.find("{http://purl.org/dc/elements/1.1/}title")
        self.assertIsNotNone(title)
        self.assertEqual(title.text, "ТКП")
        parts = app.find(
            "{http://schemas.openxmlformats.org/officeDocument/2006/extended-properties}"
            "TitlesOfParts"
        )
        assert parts is not None
        first = parts.find(
            "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}vector"
            "/{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}lpstr"
        )
        self.assertIsNotNone(first)
        self.assertEqual(first.text, "ТКП")

    def test_template_notes_who_lays_the_external_cables(self):
        """Appendix 5 says who wires the field cables and what is not included."""
        note = next(
            paragraph
            for paragraph in template_body().findall(f"{{{W}}}p")
            if element_text(paragraph).startswith("Примечание: Прокладка и расключение")
        )
        text = element_text(note)
        self.assertTrue(
            text.startswith(
                "Примечание: Прокладка и расключение внешних кабельных связей от "
                "датчиков/механизмов до шкафов управления выполняются силами Заказчика."
            ),
            text,
        )
        for excluded in (
            "кабеленесущих систем",
            "проходок",
            "огнезащита",
            "демонтаж",
        ):
            self.assertIn(excluded, text)

    def test_template_lists_site_readiness_requirements(self):
        """Appendix 5 tells the customer how the site has to be prepared.

        The block has no heading of its own: the sentence above the list
        introduces it, the way the paragraphs around it introduce theirs. The
        list is deliberately the short one — site readiness is stated once, in
        the appendix next to the works it belongs to.
        """
        body = template_body()
        paragraphs = body.findall(f"{{{W}}}p")
        texts = [element_text(p).strip() for p in paragraphs]
        self.assertNotIn("Требования к готовности площадки:", texts)
        intro = texts.index("К началу монтажных работ Заказчик обеспечивает:")
        note = next(
            index
            for index, text in enumerate(texts)
            if text.startswith("Примечание: Прокладка")
        )
        mounting = find_table(body, "Монтаж шкафов на подготовленные основания")
        # Order: the works table, then the readiness block, then the note.
        self.assertLess(body.index(mounting), body.index(paragraphs[intro]))
        self.assertLess(body.index(paragraphs[intro]), body.index(paragraphs[note]))
        block = texts[intro:note]
        for topic in (
            "основания и фальшпол",
            "контур заземления",
            "электропитание",
            "освещение",
            "пропускной режим",
            "бытовки",
            "места складирования",
        ):
            self.assertTrue(any(item.startswith(topic) for item in block), topic)
        self.assertIn("~220 В", " ".join(block))
        self.assertIn("~380 В", " ".join(block))
        # Every item of the list is one bullet of the same list.
        items = block[1:]
        self.assertEqual(len(items), 7)
        for paragraph in paragraphs[intro + 1 : note]:
            self.assertIsNotNone(paragraph.find(f"{{{W}}}pPr/{{{W}}}numPr"))

    def page_two_texts(self) -> list[str]:
        return [element_text(p).strip() for p in template_body().findall(f"{{{W}}}p")]

    def test_template_names_the_basis_of_the_offer(self):
        """Page 2 says what the offer rests on.

        The block is one sentence: the review cut the four-bullet list down to
        the document the offer is based on, because the lines about the two
        re-estimate triggers only repeated what the parameters table and
        Appendix 8 already say. The document itself stays inside that sentence —
        on a line of its own it read as a form to fill in rather than as the
        basis of the offer.
        """
        texts = self.page_two_texts()
        self.assertIn("Исходные данные и допущения:", texts)
        basis = texts.index("Исходные данные и допущения:")
        exclusions = texts.index("Не входит в объём предложения:")
        parameters = texts.index("Параметры предложения")
        self.assertLess(parameters, basis)
        self.assertLess(basis, exclusions)
        sentence = texts[basis + 1]
        self.assertTrue(
            sentence.startswith(
                "Предложение подготовлено на основании следующих исходных данных:"
            ),
            sentence,
        )
        self.assertIn("документ-основание — ____________ (наименование)", sentence)
        self.assertIn("№ ____ от __.__.____, ревизия ____;", sentence)
        # No second line and no double colon: the blank follows the same sentence.
        self.assertNotIn("данных:документ", sentence)
        self.assertNotIn(":", sentence[sentence.index(":") + 1 :])
        self.assertEqual(texts[basis + 2], "Условия оплаты:")
        # The wording the review rejected: too general to hold anybody to it.
        self.assertNotIn("могут быть уточнены", " ".join(texts[basis:exclusions]))

    def test_template_lists_what_the_offer_excludes(self):
        """Page 2 draws the boundary of the supply.

        The bullets follow the heading directly: the introduction that used to
        open them («В объём и стоимость настоящего предложения не входят:»)
        said the same thing twice.
        """
        texts = self.page_two_texts()
        self.assertNotIn("Допущения:", texts)
        self.assertNotIn(
            "В объём и стоимость настоящего предложения не входят:", texts
        )
        heading = texts.index("Не входит в объём предложения:")
        warranty = texts.index("Гарантийные обязательства:")
        self.assertLess(heading, warranty)
        block = texts[heading:warranty]
        for topic in (
            "датчиков",
            "исполнительных механизмов",
            "внешние кабельные связи",
            "строительная часть",
            "АРМ и серверов верхнего уровня",
            "ИБП",
            "стадий «П» и «Р»",
            "метрологическая аттестация",
            "экспертиза промышленной безопасности",
            "КИИ",
            "ФСТЭК России",
            "надзорных органах",
        ):
            self.assertIn(topic, " ".join(block))
        # ЗИП is not excluded here: the note after Appendix 3 already limits the
        # scope of supply to the items of the specification.
        self.assertNotIn("ЗИП", " ".join(block))

    def test_template_drops_the_customer_duties_block(self):
        """Page 2 no longer repeats what Appendix 5 asks the customer for.

        The block stood between the exclusions and the warranty and listed the
        same site readiness requirements the mounting appendix states in detail.
        The abbreviated page two keeps one place for them.
        """
        texts = self.page_two_texts()
        self.assertNotIn("Обязательства Заказчика:", texts)
        self.assertNotIn("Заказчик обеспечивает:", texts)
        # The guarantees follow the exclusions directly.
        exclusions = texts.index("Не входит в объём предложения:")
        warranty = texts.index("Гарантийные обязательства:")
        self.assertTrue(all(texts[exclusions + 1 : warranty]))

    def test_template_states_the_total_project_duration(self):
        """The parameters table sums the stages up to 35 weeks."""
        body = template_body()
        rows = rows_of(find_table(body, "ИТОГО без НДС"))
        labels = [row[1] if len(row) > 2 else row[0] for row in rows]
        self.assertIn("Общий срок реализации проекта", labels)
        duration = rows[labels.index("Общий срок реализации проекта")]
        self.assertEqual(duration[2], "35 недель")
        # It closes the list of stages, before the money.
        self.assertLess(
            labels.index("Обучение персонала"),
            labels.index("Общий срок реализации проекта"),
        )
        self.assertLess(
            labels.index("Общий срок реализации проекта"), labels.index("ИТОГО без НДС")
        )

    def test_template_ends_with_terms_and_abbreviations(self):
        """The last page explains the abbreviations used in the proposal."""
        body = template_body()
        paragraphs = body.findall(f"{{{W}}}p")
        heading = next(
            p for p in paragraphs if element_text(p).strip() == "Термины и сокращения"
        )
        # It opens a page of its own at the very end of the document: the
        # heading, the sentence above the table and the table itself, with only
        # spacer paragraphs after them (Word keeps one before the section
        # properties, so «nothing but the table follows the heading» is not a
        # property the document can have).
        self.assertIsNotNone(heading.find(f"{{{W}}}pPr/{{{W}}}pageBreakBefore"))
        tail = list(body)[list(body).index(heading) :]
        kinds = [etree.QName(child).localname for child in tail]
        self.assertEqual(kinds[0], "p")
        self.assertEqual(kinds[1], "p")
        self.assertEqual(kinds[2], "tbl")
        for child in tail[3:]:
            kind = etree.QName(child).localname
            self.assertIn(kind, {"p", "sectPr"})
            if kind == "p":
                self.assertEqual(element_text(child).strip(), "")
        terms = rows_of(body.findall(f"{{{W}}}tbl")[-1])
        self.assertEqual(terms[0], ["Сокращение", "Расшифровка"])
        self.assertEqual(
            [row[0] for row in terms[1:9]],
            ["ПТК", "ППО", "РКД", "УСО", "ЗПИ", "ПСИ", "АРМ", "КИИ"],
        )
        abbreviations = {row[0] for row in terms[1:]}
        for used in (
            "АСУ ТП", "ЗИП", "ИБП", "ИМ", "ЛВС", "НДС", "ПД", "ПМИ", "ПНР", "СИ",
            "ТО", "ФСТЭК России", "ЭД",
        ):
            self.assertIn(used, abbreviations)
        for row in terms[1:]:
            self.assertEqual(len(row), 2)
            self.assertTrue(row[1], row)

    def test_template_body_text_is_twelve_point(self):
        """10 pt and 11 pt are retired, with one deliberate exception.

        The document is set in 12 pt ("24"), and 10 pt ("20") is gone from the
        text entirely. The table of stages in Appendix 8 stays at 11 pt ("22"):
        four columns of names and terms wrap badly at 12 pt, and the table reads
        as a schedule rather than as prose, so the smaller size is the decision
        the sync step protects instead of undoing.
        """
        document = read_part("word/document.xml")
        styles = read_part("word/styles.xml")
        self.assertEqual(
            {size.get(f"{{{W}}}val") for size in document.iter(f"{{{W}}}sz")} & {"20"},
            set(),
        )
        plan = [
            table
            for table in document.iter(f"{{{W}}}tbl")
            if "Плановая дата выполнения" in element_text(table)
        ]
        self.assertEqual(len(plan), 1)
        calendar = plan[0]
        def in_calendar(size) -> bool:
            ancestor = size.getparent()
            while ancestor is not None:
                if ancestor is calendar:
                    return True
                ancestor = ancestor.getparent()
            return False

        sizes = {}
        for tag in ("sz", "szCs"):
            for size in document.iter(f"{{{W}}}{tag}"):
                value = size.get(f"{{{W}}}val")
                if value in {"20", "22"}:
                    sizes.setdefault(value, set()).add(in_calendar(size))
        # Every 11 pt run of the document belongs to that table, and the table
        # holds nothing else.
        self.assertEqual(sizes.get("22"), {True})
        self.assertEqual(
            {size.get(f"{{{W}}}val") for size in calendar.iter(f"{{{W}}}sz")}, {"22"}
        )
        defaults = styles.find(f"{{{W}}}docDefaults")
        for tag in ("sz", "szCs"):
            self.assertEqual(
                defaults.find(f".//{{{W}}}{tag}").get(f"{{{W}}}val"), "24", tag
            )
        used = {style.get(f"{{{W}}}val") for style in document.iter(f"{{{W}}}pStyle")}
        used |= {style.get(f"{{{W}}}val") for style in document.iter(f"{{{W}}}rStyle")}
        for style in styles.findall(f"{{{W}}}style"):
            if style.get(f"{{{W}}}styleId") not in used:
                continue
            own = {
                size.get(f"{{{W}}}val")
                for tag in ("sz", "szCs")
                for size in style.iter(f"{{{W}}}{tag}")
            }
            self.assertEqual(own & {"20"}, set(), style.get(f"{{{W}}}styleId"))

    def test_running_header_takes_the_number_from_the_first_page(self):
        """The running head repeats the «Номер» row of the first page."""
        header = read_part("word/header1.xml")
        paragraph = next(
            p
            for p in header.iter(f"{{{W}}}p")
            if "Технико-коммерческое предложение" in element_text(p)
        )
        starts = read_part("word/document.xml").xpath(
            ".//w:bookmarkStart[@w:name='Num']", namespaces=NS
        )
        number = element_text(starts[0].getparent())
        self.assertTrue(number.startswith("№"), number)
        self.assertEqual(
            element_text(paragraph), f"Технико-коммерческое предложение {number}"
        )
        self.assertIn(
            "REF Num",
            " ".join(i.text or "" for i in paragraph.iter(f"{{{W}}}instrText")),
        )
        # The field sits inside the HeaderNum bookmark, which both generators
        # replace with the number read from the workbook.
        children = list(paragraph)
        start = next(
            index
            for index, child in enumerate(children)
            if child.tag == f"{{{W}}}bookmarkStart"
            and child.get(f"{{{W}}}name") == "HeaderNum"
        )
        end = next(
            index
            for index, child in enumerate(children)
            if child.tag == f"{{{W}}}bookmarkEnd"
            and child.get(f"{{{W}}}id") == children[start].get(f"{{{W}}}id")
        )
        field = next(
            index
            for index, child in enumerate(children)
            if child.find(f"{{{W}}}fldChar") is not None
        )
        self.assertLess(start, field)
        self.assertLess(field, end)

    def test_number_keeps_a_single_sign(self):
        """The head reads «…предложение № 104/1…», never «№ № 104/1»."""
        self.assertEqual(format_number("104/1 от 16.10.2026"), "№ 104/1 от 16.10.2026")
        self.assertEqual(format_number("№ 104/1 от 16.10.2026"), "№ 104/1 от 16.10.2026")
        self.assertEqual(format_number("№104/1"), "№ 104/1")
        self.assertEqual(format_number("  12/3  "), "№ 12/3")
        self.assertEqual(format_number(None), "")
        self.assertEqual(format_number("#REF!"), "")

    def test_template_uses_the_reviewed_content(self):
        """Content changes from the document review: rows, order, terms, titles."""
        with ZipFile(TEMPLATE) as archive:
            document = etree.fromstring(archive.read("word/document.xml"))
            header = etree.fromstring(archive.read("word/header1.xml"))
        body = document.find(f"{{{W}}}body")
        assert body is not None
        tables = body.findall(f"{{{W}}}tbl")

        def table_rows(index: int) -> list[list[str]]:
            return rows_of(tables[index])

        # Приложение 4 is the list of documents handed over to the customer. The
        # shaded «РКД и ЭД по ГОСТ 34» row that used to open it is gone: the
        # heading of the appendix says the same, and the list starts at 1.
        documents = table_rows(2)
        self.assertEqual(
            [row[-1] for row in documents],
            [
                "Наименование",
                "Схема структурная комплекса технических средств автоматизации",
                "Схемы соединений внешних проводок и таблицы соединений внешних проводок",
                "Схемы принципиальные и схемы подключений",
                "Чертежи общего вида шкафов",
                "Ведомость покупных изделий",
                "Программа и методика испытаний (ПМИ)",
                "Протоколы заводских приёмо‑сдаточных испытаний (ЗПИ)",
                "Протоколы приёмо‑сдаточных испытаний (ПСИ)",
                "Описание прикладного программного обеспечения (ППО)",
                "Руководство оператора",
                "Руководство администратора",
                "Инструкция по техническому обслуживанию",
                "Формуляр (паспорт) на оборудование",
                "Ведомость ЗИП",
                "Протокол проверки наличия цепи между заземлителями и заземляемыми элементами",
                "Протокол измерения сопротивления изоляции",
                "Протокол проверки автоматических выключателей в электроустановках "
                "напряжением до 1000 В.",
                "Кабельный журнал",
                "Исполнительные схемы",
                "Акты скрытых работ",
                "Паспорта и сертификаты соответствия на покупное оборудование",
                "Акт выполненных пусконаладочных работ",
            ],
        )
        # Every row is a numbered document, 1..22 — there is no spanning row.
        self.assertEqual(len(documents[1]), 2)
        self.assertNotIn(
            "Оформляется на этапе выполнения монтажных работ", element_text(tables[2])
        )
        # Приложение 5: unloading, mounting, marking, power, earthing, cabling.
        self.assertEqual(
            [row[-1] for row in table_rows(3)],
            [
                "Наименование работ",
                "Разгрузка оборудования и перемещение к месту установки",
                "Монтаж шкафов на подготовленные основания",
                "Маркировка шкафов и оборудования",
                "Подключение питания шкафов",
                "Заземление шкафов",
                "Монтаж межшкафных связей",
                "Монтаж ЛВС и АРМ",
                "Сборка кроссов",
                "Маркировка кабелей",
            ],
        )
        # Приложение 6: integration with the upper level comes before the
        # comprehensive test.
        works = [r[1] for r in table_rows(4)]
        self.assertLess(
            works.index("Интеграция с верхним уровнем"), works.index("Комплексное опробование")
        )
        self.assertIn("Настройка контроллеров ПТК", works)
        # Приложение 4 is listed under the new name, and the reference pages at
        # the end are listed as the appendices 9 and 10 they now are.
        appendix_list = table_rows(1)
        listed = [" ".join(row).strip() for row in appendix_list]
        self.assertEqual(
            listed[1:],
            [
                "Перечень сигналов АСУ ТП",
                "Структура ПТК",
                "Спецификация оборудования",
                "Перечень передаваемой документации",
                "Монтажные работы",
                "Пусконаладочные работы",
                "Обучение персонала",
                "Календарный план",
                "Описание ПТК АРКС400",
                "О компании «ЭКСАРА»",
            ],
        )
        # Every entry keeps the automatic number of the appendix it names.
        numbers = [
            row.find(f".//{{{W}}}numPr/{{{W}}}numId").get(f"{{{W}}}val")
            for row in tables[1].findall(f"{{{W}}}tr")[1:]
        ]
        self.assertEqual(numbers, ["25"] * 10)
        # The document is a technical and commercial proposal, in the title and
        # in the running header.
        self.assertIn("Технико-коммерческое предложение", "".join(body.itertext()))
        self.assertIn("Технико-коммерческое предложение", "".join(header.itertext()))

    def test_template_draws_the_calendar_plan_as_a_gantt_chart(self):
        """Appendix 8 repeats its stages as bars, one column per week.

        The chart follows the table of stages directly: the appendix is the
        calendar plan, so a heading and a caption in front of the bars repeated
        the heading of the appendix itself. The table is laid out to the full
        page — wider than the text area on purpose, so that the week columns stay
        legible at 8 pt — and it is the only chart in the document.
        """
        body = template_body()
        plan = find_table(
            body, "20 (двадцать) недель с даты поступления авансового платежа"
        )
        charts = [
            table for table in body.findall(f"{{{W}}}tbl") if rows_of(table)[0][0] == "Этап"
        ]
        self.assertEqual(len(charts), 1)
        chart = charts[0]
        # A section break stands between the table of stages and its chart: the
        # chart needs the full 14574-twip grid, which does not fit a portrait
        # page, so it gets a landscape page of its own. Nothing else stands there.
        self.assertEqual(body.index(plan) + 2, body.index(chart))
        opening = body[body.index(plan) + 1]
        self.assertEqual(opening.tag, f"{{{W}}}p")
        self.assertFalse(element_text(opening).strip())
        for gone in ("Диаграмма Ганта:", "Полосы показывают недели выполнения этапов"):
            self.assertNotIn(
                gone,
                " ".join(element_text(p) for p in body.findall(f"{{{W}}}p")),
            )

        grid = chart.find(f"{{{W}}}tblGrid").findall(f"{{{W}}}gridCol")
        widths = [int(column.get(f"{{{W}}}w")) for column in grid]
        # One label column and 35 week columns — no spare column at the right.
        self.assertEqual(len(grid), 36)
        self.assertEqual(widths[0], 4949)
        self.assertEqual(set(widths[1:]), {275})
        self.assertEqual(sum(widths), 4949 + 35 * 275)

        rows = rows_of(chart)
        self.assertEqual(rows[0][0], "Этап")
        self.assertEqual(rows[0][1:], [str(week) for week in range(1, 36)])
        stage_rows = chart.findall(f"{{{W}}}tr")
        for row in stage_rows:
            span = sum(
                int(cell.find(f"{{{W}}}tcPr/{{{W}}}gridSpan").get(f"{{{W}}}val"))
                if cell.find(f"{{{W}}}tcPr/{{{W}}}gridSpan") is not None
                else 1
                for cell in row.findall(f"{{{W}}}tc")
            )
            self.assertEqual(span, len(grid), element_text(row)[:40])

        # The bars follow the critical path: equipment, factory tests, delivery,
        # mounting and commissioning add up to the 35 weeks of page two.
        bars = {}
        for index in range(1, len(stage_rows) - 1):
            weeks = [
                week
                for week, cell in enumerate(
                    stage_rows[index].findall(f"{{{W}}}tc")[1:], start=1
                )
                if cell.find(f"{{{W}}}tcPr/{{{W}}}shd") is not None
            ]
            bars[rows[index][0]] = (weeks[0], weeks[-1], len(weeks))
        self.assertEqual(
            bars,
            {
                "Разработка РКД": (1, 10, 10),
                "Разработка ППО": (1, 19, 19),
                "Заказ и сборка оборудования": (1, 20, 20),
                "Заводские приёмо‑сдаточные испытания": (21, 21, 1),
                "Доставка оборудования": (22, 22, 1),
                "Монтажные работы": (23, 25, 3),
                "Пусконаладочные работы": (26, 35, 10),
                "Обучение персонала": (35, 35, 1),
            },
        )
        self.assertEqual(rows[-1], ["Общий срок реализации проекта", "35 недель"])
        total = stage_rows[-1].findall(f"{{{W}}}tc")
        self.assertEqual(
            total[1].find(f"{{{W}}}tcPr/{{{W}}}gridSpan").get(f"{{{W}}}val"), "35"
        )

    def test_gantt_chart_is_alone_on_a_landscape_page(self):
        """The chart is the only thing on a page with the sheet turned.

        The grid is wider than the portrait text area by design, so the appendix
        hands the chart a section of its own: a break before the chart and one
        after it, and the rest of the document stays portrait.
        """
        body = template_body()
        plan = find_table(
            body, "20 (двадцать) недель с даты поступления авансового платежа"
        )
        chart = next(
            table for table in body.findall(f"{{{W}}}tbl") if rows_of(table)[0][0] == "Этап"
        )
        opening = body[body.index(plan) + 1]
        closing = body[body.index(chart) + 1]

        def section_of(paragraph) -> etree._Element:
            found = paragraph.find(f"{{{W}}}pPr/{{{W}}}sectPr")
            self.assertIsNotNone(found, element_text(paragraph) or "no section break")
            return found

        def orientation(sectPr) -> tuple[str, bool]:
            page = sectPr.find(f"{{{W}}}pgSz")
            return page.get(f"{{{W}}}w"), page.get(f"{{{W}}}orient") == "landscape"

        # The break before the chart closes the portrait page of the appendix.
        self.assertFalse(orientation(section_of(opening))[1])
        # The break after the chart closes the landscape page it sits on.
        self.assertTrue(orientation(section_of(closing))[1])
        # The chart is the whole of that section: nothing but the break follows
        # it, and nothing but the break precedes it.
        self.assertEqual(body.index(chart), body.index(opening) + 1)
        self.assertEqual(body.index(closing), body.index(chart) + 1)
        # Landscape A4 with the margins of the document holds the full grid.
        page = section_of(closing).find(f"{{{W}}}pgSz")
        margins = section_of(closing).find(f"{{{W}}}pgMar")
        text_area = int(page.get(f"{{{W}}}w")) - int(
            margins.get(f"{{{W}}}left")
        ) - int(margins.get(f"{{{W}}}right"))
        self.assertEqual(int(page.get(f"{{{W}}}h")), 11906)
        self.assertGreaterEqual(text_area, 4949 + 35 * 275)
        # The chart does not turn the rest of the document with it: the next
        # section break is a portrait page again (Appendix 9), and so is the
        # last section (the list of terms).
        elements = list(body)
        following = [
            element.find(f"{{{W}}}pPr/{{{W}}}sectPr")
            for element in elements[elements.index(closing) + 1 :]
            if element.tag == f"{{{W}}}p"
        ]
        nxt = next(sectPr for sectPr in following if sectPr is not None)
        self.assertFalse(orientation(nxt)[1])
        self.assertFalse(
            orientation(body.find(f"{{{W}}}sectPr"))[1], "the last section is portrait"
        )

    def test_template_numbers_the_reference_pages_as_appendices(self):
        """АРКС400 and the company page are appendices 9 and 10, numbered by field."""
        paragraphs = template_body().findall(f"{{{W}}}p")
        labels = [
            paragraph
            for paragraph in paragraphs
            if any(
                "SEQ App" in (instruction.text or "")
                for instruction in paragraph.iter(f"{{{W}}}instrText")
            )
        ]
        self.assertEqual(
            [element_text(label).strip() for label in labels],
            [f"Приложение {number}" for number in range(1, 11)],
        )
        # Each label opens its own page, so the title follows it immediately.
        titles = [element_text(label.getnext()).strip() for label in labels]
        self.assertTrue(titles[8].startswith("АРКС400 — современные"), titles[8])
        self.assertTrue(titles[9].startswith("О компании «ЭКСАРА»"), titles[9])
        # The terms page stays outside the appendices: no label, no SEQ field.
        terms = next(
            paragraph
            for paragraph in paragraphs
            if element_text(paragraph).strip() == "Термины и сокращения"
        )
        self.assertNotIn(terms, [label.getnext() for label in labels])
        self.assertFalse(
            any(
                "SEQ App" in (instruction.text or "")
                for instruction in terms.iter(f"{{{W}}}instrText")
            )
        )

    def test_preview_has_only_the_standard_specification_columns(self):
        # Read the sheet by name: "xl/worksheets/sheet26.xml" would silently point
        # at an unrelated worksheet after any insert, removal or reordering.
        wb = openpyxl.load_workbook(WORKBOOK)
        self.assertIn("Спец_ТКП", wb.sheetnames)
        preview = wb["Спец_ТКП"]
        self.assertEqual(preview["A1"].value, "№ п/п")
        self.assertTrue(preview.column_dimensions["F"].hidden)

        source = (ROOT / "Module1.bas").read_text(encoding="utf-8")
        self.assertNotIn('"Тип"', source)
        self.assertIn("For col = 1 To 5", source)  # Word receives visible A:E verbatim.

    def test_generation_groups_sums_skips_and_refreshes(self):
        app = ExcelApplication.open(WORKBOOK)
        rows = [
            ("SW(w.o.VAT)", 14, "Лицензия ПО", "MasterPLC.N1000", 2),
            ("1", 3, None, "NV-1", 2),  # broken lookup in C: fall back to article D
            ("1", 14, "Шкаф Эксара", "EX-1", 4),
            ("1", 20, "Модуль ввода", "X-123", 2),
            ("2", 20, "Модуль ввода", "X-123", 3),
            ("1", 21, "Не включать", "X-0", 0),
            ("Spare parts", 25, "Реле ЗИП", "ZIP-1", 1),
        ]
        for sheet, row, description, article, qty in rows:
            source = app.sheet(sheet)
            if description is not None:
                source.set_value(f"C{row}", description)
            source.set_value(f"D{row}", article)
            source.set_value(f"G{row}", qty)

        app.run("ОбновитьСпецификациюТКП")
        preview = app.sheet("Спец_ТКП")
        values = [
            tuple(preview.value(f"{letter}{row}") for letter in "ABCDEF")
            for row in range(2, 14)
        ]
        group_rows = [v for v in values if v[5] == "H"]
        item_rows = [v for v in values if v[5] == "I"]

        self.assertEqual(
            [v[1] for v in group_rows],
            [
                "Программное обеспечение",
                "Оборудование ПТК (шкаф 1)",
                "Телекоммуникационное оборудование для АСУ ТП (шкаф 1)",
                "Шкафное оборудование ПТК (шкаф 1)",
                "Шкафное оборудование ПТК (шкаф 2)",
                "ЗИП",
            ],
        )
        self.assertEqual([v[0] for v in group_rows], ["1.", "2.", "3.", "4.", "5.", "6."])
        self.assertEqual([v[3] for v in group_rows], ["Компл."] * 6)
        self.assertEqual([v[4] for v in group_rows], [1.0] * 6)

        self.assertEqual(
            [v[1] for v in item_rows],
            [
                "Лицензия ПО",
                "NV-1",
                "Шкаф Эксара",
                "Модуль ввода",
                "Модуль ввода",
                "Реле ЗИП",
            ],
        )
        self.assertEqual([v[0] for v in item_rows], ["1."] * 6)
        self.assertEqual(preview.value("A1"), "№ п/п")
        self.assertEqual(preview.value("A14"), None)

        # Verify Python generator produces the exact same rows on the same input data
        wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
        for sheet, row, description, article, qty in rows:
            ws = wb[sheet]
            if description is not None:
                ws[f"C{row}"] = description
            ws[f"D{row}"] = article
            ws[f"G{row}"] = qty
        py_data = read_tkp_data(wb)
        py_rows = build_preview_rows(py_data.sections)
        self.assertEqual(
            normalise_preview_rows(py_rows), normalise_preview_rows(values)
        )

        # A second run must clear old rows: the preview is regenerated, not appended.
        for sheet, row, *_ in rows:
            app.sheet(sheet).set_value(f"G{row}", 0)
        app.run("ОбновитьСпецификациюТКП")
        self.assertEqual(preview.value("A1"), "№ п/п")
        self.assertEqual(preview.value("A2"), None)
        self.assertEqual(preview.value("B13"), None)

    def _run_both(self, edits, last_row=14):
        """Apply `edits` to both engines and return their normalised preview rows.

        Each edit is (sheet, row, description, article, quantity) written into the
        C/D/G columns. Neither engine writes to disk: pyopenvba interprets the VBA
        in memory and openpyxl only mutates the loaded workbook.
        """
        app = ExcelApplication.open(WORKBOOK)
        for sheet, row, description, article, qty in edits:
            source = app.sheet(sheet)
            source.set_value(f"C{row}", description)
            source.set_value(f"D{row}", article)
            source.set_value(f"G{row}", qty)
        app.run("ОбновитьСпецификациюТКП")
        preview = app.sheet("Спец_ТКП")
        values = [
            tuple(preview.value(f"{letter}{row}") for letter in "ABCDEF")
            for row in range(2, last_row)
        ]
        vba_rows = normalise_preview_rows([v for v in values if v[1] is not None])

        wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
        for sheet, row, description, article, qty in edits:
            ws = wb[sheet]
            ws[f"C{row}"] = description
            ws[f"D{row}"] = article
            ws[f"G{row}"] = qty
        py_rows = normalise_preview_rows(build_preview_rows(read_tkp_data(wb).sections))
        return vba_rows, py_rows

    def test_total_and_margin_rows_are_skipped_in_both_engines(self):
        """«Итого»/«МАРЖА» in the description column must never reach Appendix 3.

        Regression (H1): the VBA implementation inspected the caption column (B)
        only, so a total typed into column C was emitted as a normal line item.
        """
        # A total ends its block, so the ordinary positions are placed before it,
        # matching the real sheet layout (rows 3-9, then «Итого…» on row 10).
        vba_rows, py_rows = self._run_both(
            [
                ("1", 4, "Модуль ввода", "ART-1", 2),
                ("1", 5, "Обычная позиция", "ART-2", 3),
                ("1", 6, "Итого по шкафу", "ART-TOTAL", 5),
                ("1", 7, "МАРЖА по шкафу", "ART-MARGIN", 7),
            ]
        )
        self.assertEqual(vba_rows, py_rows)
        descriptions = [row[1] for row in py_rows]
        self.assertIn("Оборудование ПТК (шкаф 1)", descriptions)
        self.assertIn("Модуль ввода", descriptions)
        self.assertIn("Обычная позиция", descriptions)
        self.assertNotIn("Итого по шкафу", descriptions)
        self.assertNotIn("МАРЖА по шкафу", descriptions)

    def test_hash_prefixed_values_survive_in_both_engines(self):
        """A value starting with "#" is data, not an Excel error.

        Regression (H2): clean_text() blanked every value beginning with "#", so
        an article such as "#A-100" was silently dropped from the specification.
        """
        vba_rows, py_rows = self._run_both(
            [
                ("1", 4, "Контроллер", "#A-100", 3),
                ("1", 5, "#1 Модуль", "ART-1", 2),
                ("1", 6, "Позиция", "ART-2", 4),
            ]
        )
        self.assertEqual(vba_rows, py_rows)
        self.assertIn("#A-100", [row[2] for row in py_rows])
        self.assertIn("#1 Модуль", [row[1] for row in py_rows])

    def test_excel_error_literals_are_still_blanked(self):
        """The error set must keep real Excel errors out of the document."""
        for error in ("#REF!", "#N/A", "#VALUE!", "#DIV/0!", "#NUM!"):
            self.assertEqual(clean_text(error), "")

    def test_format_qty_is_always_plain_decimal(self):
        """Quantities must never reach the document in scientific notation.

        Regression (M1): %g renders 1234567.5 as "1.23457e+06" and truncates to
        six significant digits.
        """
        self.assertEqual(format_qty(1), "1")
        self.assertEqual(format_qty(50), "50")
        self.assertEqual(format_qty(2.5), "2.5")
        self.assertEqual(format_qty(0.3), "0.3")
        self.assertEqual(format_qty(1234567.5), "1234567.5")
        self.assertEqual(format_qty(10000000.5), "10000000.5")
        for value in (2.5, 1 / 3, 1234567.5, 10000000.5, 0.1 + 0.2):
            self.assertNotIn("e", format_qty(value).lower())

    def test_circular_reference_resolves_to_empty(self):
        """A reference cycle degrades to a blank value instead of crashing.

        Regression (M3): _eval_cell() recursed without a guard, so two cells that
        point at each other raised an unhandled RecursionError.
        """
        wb = openpyxl.Workbook()
        wb.active.title = "Наценки-Скидки"
        first = wb.create_sheet("X")
        second = wb.create_sheet("Y")
        first["C1"] = "=Y!C1"
        second["C1"] = "=X!C1"
        self.assertEqual(_eval_cell(first, first, wb, wb, 1, 3), "")

        # A ordinary reference chain must still resolve.
        first["C1"] = "=Наценки-Скидки!C2"
        wb["Наценки-Скидки"]["C2"] = "Значение"
        self.assertEqual(_eval_cell(first, first, wb, wb, 1, 3), "Значение")

    def test_python_build_from_dotx_template(self):
        with TemporaryDirectory() as tmp:
            out_path = Path(tmp) / "out.docx"
            data = generate_tkp(
                excel_source=WORKBOOK,
                output_path=out_path,
                use_demo_if_empty=True,
            )
            self.assertTrue(out_path.exists())
            self.assertEqual(len(data.sections), 7)
            self.assertEqual(data.total_items, 16)

            doc = docx.Document(str(out_path))
            self.assertGreaterEqual(len(doc.tables), 3)
            spec_tbl = doc.tables[2]
            self.assertEqual(
                [c.text.strip() for c in spec_tbl.rows[0].cells],
                [
                    "№ п/п",
                    "Наименование и техническая характеристика",
                    "Артикул",
                    "Ед. изм.",
                    "Кол-во",
                ],
            )
            self.assertEqual(
                [c.text.strip() for c in spec_tbl.rows[1].cells], ["1", "2", "3", "4", "5"]
            )
            self.assertEqual(
                [c.text.strip() for c in spec_tbl.rows[2].cells],
                ["1.", "Программное обеспечение", "", "Компл.", "1"],
            )
            self.assertTrue(all(c.paragraphs[0].runs[0].bold for c in spec_tbl.rows[2].cells))
            self.assertFalse(any(c.paragraphs[0].runs[0].bold for c in spec_tbl.rows[3].cells))
            # The specification table follows the 12 pt house size.
            self.assertEqual(
                spec_tbl.rows[3].cells[1].paragraphs[0].runs[0].font.size, Pt(12)
            )
            # The running head carries the number of the first page, once.
            head = "".join(
                t.text or "" for t in doc.sections[0].header._element.iter(f"{{{W}}}t")
            )
            self.assertIn(
                "Технико-коммерческое предложение № 104/1 от 16.10.2026", head
            )
            self.assertNotIn("№ №", head)
            # The field of the template is gone: the head holds the sheet value.
            self.assertEqual(
                list(doc.sections[0].header._element.iter(f"{{{W}}}instrText")), []
            )


if __name__ == "__main__":
    unittest.main()
