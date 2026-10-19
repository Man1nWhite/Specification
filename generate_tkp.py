#!/usr/bin/env python3
"""Generate the commercial proposal Word document (Appendix 3 & title attributes)
directly from Gen_TKP.xlsx / Gen_TKP.xlsm using openpyxl and python-docx.

Can be used both as a library module and as a CLI script without Excel/Word COM.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from dataclasses import dataclass, field
from io import BytesIO
import math
from pathlib import Path
import re
from typing import Iterable
import warnings
from zipfile import ZipFile, ZIP_DEFLATED

import docx
from docx.document import Document as DocxDocument
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from lxml import etree
import openpyxl
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parent
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
W = f"{{{W_NS}}}"
NS = {"w": W_NS}

TEMPLATE_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
)
DOCUMENT_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
)

SPEC_HEADERS = (
    "№ п/п",
    "Наименование и техническая характеристика",
    "Артикул",
    "Ед. изм.",
    "Кол-во",
)

# Excel error literals. These are blanked out wherever they appear; any other
# value starting with "#" (a legitimate article such as "#A-100") is kept.
EXCEL_ERRORS = frozenset(
    {
        "#REF!",
        "#N/A",
        "#VALUE!",
        "#DIV/0!",
        "#NAME?",
        "#NULL!",
        "#NUM!",
        "#SPILL!",
        "#CALC!",
        "#GETTING_DATA",
    }
)


@dataclass
class SpecItem:
    description: str
    article: str
    unit: str
    qty: float


@dataclass
class SpecSection:
    title: str
    items: list[SpecItem] = field(default_factory=list)


@dataclass
class TKPMetadata:
    number: str = ""
    customer: str = ""
    object_name: str = ""
    subject: str = ""


@dataclass
class TKPData:
    metadata: TKPMetadata
    sections: list[SpecSection]

    @property
    def total_items(self) -> int:
        return sum(len(s.items) for s in self.sections)


def clean_text(value: object) -> str:
    """Normalize cell values into a single-line trimmed string, ignoring errors/None."""
    if value is None or isinstance(value, bool):
        return ""
    text = str(value)
    for ch in ("\r\n", "\r", "\n", "\t"):
        text = text.replace(ch, " ")
    text = text.strip()
    if text in EXCEL_ERRORS:
        return ""
    return text


def format_qty(qty: float) -> str:
    """Format a quantity as plain decimal notation, without trailing zeros.

    The document is a customer-facing proposal, so scientific notation must
    never reach it: %g would render 1234567.5 as "1.23457e+06".
    """
    if abs(qty - round(qty)) < 1e-9:
        return str(int(round(qty)))
    return f"{qty:.10f}".rstrip("0").rstrip(".")


def _find_sheet(wb: Workbook, *names: str) -> Worksheet | None:
    for name in names:
        if name in wb.sheetnames:
            return wb[name]
    return None


def _block_from_caption(caption: str) -> int:
    low = caption.lower()
    if "нвт" in low:
        return 1
    if "эксара" in low or "инсат" in low:
        return 2
    if "3-х поставщиков" in low:
        return 3
    return 0


def _block_from_row(row: int) -> int:
    if 3 <= row <= 9:
        return 1
    if 14 <= row <= 15:
        return 2
    if 20 <= row <= 47:
        return 3
    return 0


_SHEET_REF_RE = re.compile(r"^='?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)$")


def _eval_cell(
    ws_val: Worksheet | None,
    ws_fml: Worksheet | None,
    wb_val: Workbook | None,
    wb_fml: Workbook | None,
    row: int,
    col: int,
    _seen: frozenset[tuple[str, int, int]] = frozenset(),
) -> object:
    """Evaluate a cell from either a dual (values, formulas) workbook pair or single workbook.

    ``_seen`` holds the (sheet, row, column) keys currently being resolved. A cell
    that is already on the stack resolves to "" instead of recursing forever, so a
    circular reference degrades to an empty value rather than a RecursionError.
    """
    if ws_fml is not None:
        key = (ws_fml.title, row, col)
        if key in _seen:
            return ""
        _seen = _seen | {key}

    raw_fml = ws_fml.cell(row, col).value if ws_fml is not None else None
    raw_val = ws_val.cell(row, col).value if ws_val is not None else None

    if not (isinstance(raw_fml, str) and raw_fml.startswith("=")):
        return raw_fml if ws_fml is not None else raw_val

    upper_fml = raw_fml.upper()
    if "#REF!" in upper_fml:
        return ""

    match = _SHEET_REF_RE.match(raw_fml.strip())
    if match and wb_fml is not None:
        target_sheet, col_letters, row_digits = match.groups()
        if target_sheet in wb_fml.sheetnames:
            t_fml = wb_fml[target_sheet]
            t_val = wb_val[target_sheet] if (wb_val is not None and target_sheet in wb_val.sheetnames) else None
            col_idx = openpyxl.utils.column_index_from_string(col_letters)
            resolved = _eval_cell(
                t_val, t_fml, wb_val, wb_fml, int(row_digits), col_idx, _seen
            )
            if resolved is not None and not (isinstance(resolved, str) and resolved in EXCEL_ERRORS):
                return resolved

    if "SUMIF(" in upper_fml and ws_fml is not None and wb_fml is not None:
        article = clean_text(_eval_cell(ws_val, ws_fml, wb_val, wb_fml, row, 4, _seen))
        if not article:
            return 0
        total_qty = 0.0
        for cab_no in range(1, 18):
            sname = str(cab_no)
            if sname not in wb_fml.sheetnames:
                continue
            c_fml = wb_fml[sname]
            c_val = wb_val[sname] if (wb_val is not None and sname in wb_val.sheetnames) else None
            for r in range(3, min(c_fml.max_row + 1, 101)):
                d_val = clean_text(_eval_cell(c_val, c_fml, wb_val, wb_fml, r, 4, _seen))
                if d_val.lower() == article.lower():
                    g_val = _eval_cell(c_val, c_fml, wb_val, wb_fml, r, 7, _seen)
                    if isinstance(g_val, (int, float)) and not isinstance(g_val, bool):
                        total_qty += float(g_val)
        return math.ceil(total_qty / 10.0) if total_qty > 0 else 0

    return raw_val


def _parse_numeric_qty(raw_qty: object) -> float:
    if raw_qty is None or isinstance(raw_qty, bool):
        return 0.0
    if isinstance(raw_qty, (int, float)):
        return float(raw_qty) if not math.isnan(float(raw_qty)) else 0.0
    if isinstance(raw_qty, str):
        text = raw_qty.strip().replace(",", ".")
        if not text or text in EXCEL_ERRORS or text.startswith("="):
            return 0.0
        try:
            return float(text)
        except ValueError:
            return 0.0
    return 0.0


def _collect_section_items(
    ws_val: Worksheet | None,
    ws_fml: Worksheet,
    wb_val: Workbook | None,
    wb_fml: Workbook,
    target_block: int,
) -> list[SpecItem]:
    max_row = max(ws_fml.max_row or 1, ws_val.max_row if ws_val is not None else 1)
    current_block = 0
    has_captions = False
    items: list[SpecItem] = []

    for r in range(1, max_row + 1):
        caption = clean_text(_eval_cell(ws_val, ws_fml, wb_val, wb_fml, r, 2))
        description = clean_text(_eval_cell(ws_val, ws_fml, wb_val, wb_fml, r, 3))

        if caption.lower().startswith("спецификация "):
            current_block = _block_from_caption(caption)
            has_captions = True
            continue
        if (
            caption.lower().startswith("итого")
            or description.lower().startswith("итого")
            or description.lower().startswith("маржа")
        ):
            current_block = 0
            continue
        if caption.lower() == "№ п/п":
            continue

        eff_block = current_block
        if eff_block == 0 and not has_captions:
            eff_block = _block_from_row(r)

        if eff_block > 0 and (target_block == 0 or eff_block == target_block):
            raw_qty = _eval_cell(ws_val, ws_fml, wb_val, wb_fml, r, 7)
            qty = _parse_numeric_qty(raw_qty)
            if qty > 0:
                article = clean_text(_eval_cell(ws_val, ws_fml, wb_val, wb_fml, r, 4))
                unit = clean_text(_eval_cell(ws_val, ws_fml, wb_val, wb_fml, r, 6))
                # clean_text() already blanked Excel errors, so only an empty
                # description or a literal "0" needs the article fallback.
                if not description or description == "0":
                    description = article
                if not unit:
                    unit = "шт."
                if description:
                    items.append(
                        SpecItem(
                            description=description,
                            article=article,
                            unit=unit,
                            qty=qty,
                        )
                    )
    return items


def read_tkp_data(excel_source: str | Path | Workbook) -> TKPData:
    """Read title metadata and Appendix 3 specification sections from Excel."""
    if isinstance(excel_source, Workbook):
        wb_fml = excel_source
        wb_val = excel_source
    else:
        path = Path(excel_source)
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Data Validation extension is not supported")
            wb_val = openpyxl.load_workbook(path, data_only=True)
            wb_fml = openpyxl.load_workbook(path, data_only=False)

    info_fml = _find_sheet(wb_fml, "Наценки-Скидки")
    info_val = _find_sheet(wb_val, "Наценки-Скидки") if wb_val is not None else None
    metadata = TKPMetadata()
    if info_fml is not None:
        metadata.number = clean_text(_eval_cell(info_val, info_fml, wb_val, wb_fml, 2, 3))
        metadata.customer = clean_text(_eval_cell(info_val, info_fml, wb_val, wb_fml, 3, 3))
        metadata.object_name = clean_text(_eval_cell(info_val, info_fml, wb_val, wb_fml, 4, 3))
        metadata.subject = clean_text(_eval_cell(info_val, info_fml, wb_val, wb_fml, 5, 3))

    sections: list[SpecSection] = []

    # 1. Программное обеспечение — all positions with G > 0 from SW(w.o.VAT)
    sw_fml = _find_sheet(wb_fml, "SW(w.o.VAT)", "ПО")
    sw_val = _find_sheet(wb_val, "SW(w.o.VAT)", "ПО") if wb_val is not None else None
    if sw_fml is not None:
        sw_items = _collect_section_items(sw_val, sw_fml, wb_val, wb_fml, target_block=0)
        if sw_items:
            sections.append(SpecSection(title="Программное обеспечение", items=sw_items))

    # 2. Cabinets 1..17 — 3 blocks per cabinet sheet
    block_defs = (
        (1, "ПТК"),
        (2, "телекоммуникационное оборудование"),
        (3, "шкафное оборудование"),
    )
    for cab_no in range(1, 18):
        cab_fml = _find_sheet(wb_fml, str(cab_no), f"Шкаф {cab_no}")
        cab_val = (
            _find_sheet(wb_val, str(cab_no), f"Шкаф {cab_no}") if wb_val is not None else None
        )
        if cab_fml is None:
            continue
        for block_no, suffix in block_defs:
            cab_items = _collect_section_items(
                cab_val, cab_fml, wb_val, wb_fml, target_block=block_no
            )
            if cab_items:
                sections.append(
                    SpecSection(title=f"Шкаф {cab_no} {suffix}", items=cab_items)
                )

    # 3. ЗИП — all positions with G > 0 from Spare parts
    sp_fml = _find_sheet(wb_fml, "Spare parts", "ЗИП")
    sp_val = _find_sheet(wb_val, "Spare parts", "ЗИП") if wb_val is not None else None
    if sp_fml is not None:
        sp_items = _collect_section_items(sp_val, sp_fml, wb_val, wb_fml, target_block=0)
        if sp_items:
            sections.append(SpecSection(title="ЗИП", items=sp_items))

    return TKPData(metadata=metadata, sections=sections)


def build_preview_rows(sections: Iterable[SpecSection]) -> list[tuple[str, str, str, str, float | int, str]]:
    """Build the preview table rows matching sheet 'Спец_ТКП' (columns A..F)."""
    rows: list[tuple[str, str, str, str, float | int, str]] = []
    for sec_idx, section in enumerate(sections, start=1):
        if not section.items:
            continue
        rows.append((f"{sec_idx}.", section.title, "", "Компл.", 1, "H"))
        for item_idx, item in enumerate(section.items, start=1):
            qty_val: float | int = (
                int(round(item.qty)) if abs(item.qty - round(item.qty)) < 1e-9 else item.qty
            )
            rows.append(
                (f"{item_idx}.", item.description, item.article, item.unit, qty_val, "I")
            )
    return rows


def update_preview_sheet(wb: Workbook, sections: Iterable[SpecSection]) -> Worksheet:
    """Populate the 'Спец_ТКП' preview worksheet in an openpyxl Workbook."""
    ws = _find_sheet(wb, "Спец_ТКП")
    if ws is None:
        ws = wb.create_sheet("Спец_ТКП")

    max_r = max(ws.max_row or 1, 1)
    for r in range(1, max_r + 1):
        for c in range(1, 7):
            cell = ws.cell(r, c)
            cell.value = None
            cell.font = openpyxl.styles.Font(bold=False)

    bold_font = openpyxl.styles.Font(bold=True)
    regular_font = openpyxl.styles.Font(bold=False)
    for col_idx, header_text in enumerate(SPEC_HEADERS, start=1):
        cell = ws.cell(1, col_idx, header_text)
        cell.font = bold_font

    for row_offset, row_data in enumerate(build_preview_rows(sections), start=2):
        is_group = row_data[5] == "H"
        row_font = bold_font if is_group else regular_font
        for col_idx, val in enumerate(row_data, start=1):
            c = ws.cell(row_offset, col_idx, val)
            if col_idx <= 5:
                c.font = row_font

    ws.column_dimensions["A"].width = 7
    ws.column_dimensions["B"].width = 65
    ws.column_dimensions["C"].width = 27
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].hidden = True
    return ws


def open_word_template(template_path: str | Path) -> DocxDocument:
    """Open either a .docx or .dotx Word template in python-docx."""
    path = Path(template_path)
    buf = BytesIO()
    with ZipFile(path, "r") as zin, ZipFile(buf, "w", ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            if info.filename == "word/_rels/settings.xml.rels":
                continue
            data = zin.read(info.filename)
            if info.filename == "[Content_Types].xml":
                text = data.decode("utf-8").replace(
                    TEMPLATE_CONTENT_TYPE, DOCUMENT_CONTENT_TYPE
                )
                data = text.encode("utf-8")
            elif info.filename == "word/settings.xml":
                root = etree.fromstring(data)
                for attached in root.findall(f"{W}attachedTemplate"):
                    root.remove(attached)
                data = etree.tostring(
                    root, encoding="UTF-8", xml_declaration=True, standalone=True
                )
            zout.writestr(info, data)
    buf.seek(0)
    return docx.Document(buf)


def set_bookmark_text(root: etree._Element, name: str, text: str) -> bool:
    """Set the text inside a Word bookmark while preserving paragraph/run formatting."""
    starts = [
        el
        for el in root.findall(f".//{W}bookmarkStart")
        if el.get(f"{W}name") == name
    ]
    if not starts:
        return False

    start = starts[0]
    bm_id = start.get(f"{W}id")
    parent = start.getparent()
    if parent is None:
        return False

    ends = [
        el
        for el in root.findall(f".//{W}bookmarkEnd")
        if el.get(f"{W}id") == bm_id
    ]
    end = ends[0] if ends else None

    rpr_template = None
    ppr = parent.find(f"{W}pPr")
    if ppr is not None:
        p_rpr = ppr.find(f"{W}rPr")
        if p_rpr is not None:
            rpr_template = deepcopy(p_rpr)

    if end is not None and end.getparent() is parent:
        start_idx = list(parent).index(start)
        end_idx = list(parent).index(end)
        for child in list(parent)[start_idx + 1 : end_idx]:
            if rpr_template is None and child.tag == f"{W}r":
                c_rpr = child.find(f"{W}rPr")
                if c_rpr is not None:
                    rpr_template = deepcopy(c_rpr)
            parent.remove(child)
    else:
        if end is not None and end.getparent() is not None:
            end.getparent().remove(end)
        end = etree.Element(f"{W}bookmarkEnd", {f"{W}id": bm_id or "0"})
        parent.insert(list(parent).index(start) + 1, end)

    new_run = etree.Element(f"{W}r")
    if rpr_template is not None:
        new_run.append(rpr_template)
    t_el = etree.SubElement(new_run, f"{W}t")
    if text.startswith(" ") or text.endswith(" "):
        t_el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t_el.text = text
    parent.insert(list(parent).index(end), new_run)
    return True


def _style_cell(
    cell,
    text: str,
    *,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    bg_fill: str | None = None,
    text_color: RGBColor = RGBColor(0x00, 0x00, 0x00),
    width_pct: int | None = None,
) -> None:
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()

    if width_pct is not None:
        for old_w in tcPr.findall(qn("w:tcW")):
            tcPr.remove(old_w)
        tcW = OxmlElement("w:tcW")
        tcW.set(qn("w:w"), str(width_pct))
        tcW.set(qn("w:type"), "pct")
        tcPr.append(tcW)

    if bg_fill is not None:
        for old_shd in tcPr.findall(qn("w:shd")):
            tcPr.remove(old_shd)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{bg_fill}"/>')
        tcPr.append(shd)

    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)

    for r in list(p.runs):
        p._p.remove(r._r)

    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = text_color


def _set_row_repeat_header(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:tblHeader")) is None:
        trPr.append(OxmlElement("w:tblHeader"))


def populate_spec_table(doc: DocxDocument, sections: Iterable[SpecSection]) -> docx.table.Table:
    """Create and populate the Appendix 3 specification table at bookmark Spec_Table."""
    preview_rows = build_preview_rows(sections)
    total_rows = 2 + len(preview_rows)

    table = doc.add_table(rows=total_rows, cols=5)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    tblPr = table._tbl.tblPr
    for old_w in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old_w)
    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), "5000")
    tblW.set(qn("w:type"), "pct")
    tblPr.append(tblW)

    for old_borders in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old_borders)
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        '<w:top w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:left w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:bottom w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:right w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:insideH w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        '<w:insideV w:val="single" w:sz="4" w:space="0" w:color="BFBFBF"/>'
        "</w:tblBorders>"
    )
    tblPr.append(tblBorders)

    col_pcts = (359, 2739, 1009, 505, 388)
    grid_twips = (706, 5387, 1985, 993, 763)
    for old_grid in table._tbl.findall(qn("w:tblGrid")):
        table._tbl.remove(old_grid)
    tblGrid = OxmlElement("w:tblGrid")
    for tw in grid_twips:
        gridCol = OxmlElement("w:gridCol")
        gridCol.set(qn("w:w"), str(tw))
        tblGrid.append(gridCol)
    table._tbl.insert(list(table._tbl).index(tblPr) + 1, tblGrid)

    # Row 0: column titles, Row 1: column numbers 1..5
    _set_row_repeat_header(table.rows[0])
    _set_row_repeat_header(table.rows[1])
    for col_idx, header_text in enumerate(SPEC_HEADERS):
        _style_cell(
            table.rows[0].cells[col_idx],
            header_text,
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bg_fill="C00000",
            text_color=RGBColor(0xFF, 0xFF, 0xFF),
            width_pct=col_pcts[col_idx],
        )
        _style_cell(
            table.rows[1].cells[col_idx],
            str(col_idx + 1),
            bold=True,
            align=WD_ALIGN_PARAGRAPH.CENTER,
            bg_fill="C00000",
            text_color=RGBColor(0xFF, 0xFF, 0xFF),
            width_pct=col_pcts[col_idx],
        )

    # Data rows (sections and items)
    for idx, (num_str, desc, art, unit, qty, marker) in enumerate(preview_rows, start=2):
        row = table.rows[idx]
        is_group = marker == "H"
        bg = "D9D9D9" if is_group else "FFFFFF"
        values = (num_str, desc, art, unit, format_qty(qty))
        aligns = (
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER if is_group else WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.CENTER,
            WD_ALIGN_PARAGRAPH.CENTER,
        )
        for col_idx in range(5):
            _style_cell(
                row.cells[col_idx],
                values[col_idx],
                bold=is_group,
                align=aligns[col_idx],
                bg_fill=bg,
                width_pct=col_pcts[col_idx],
            )
        if is_group:
            for p in row.cells[0]._tc.getparent().findall(f".//{W}p"):
                pPr = p.find(f"{W}pPr")
                if pPr is not None and pPr.find(f"{W}rPr") is None:
                    rPr = etree.SubElement(pPr, f"{W}rPr")
                    etree.SubElement(rPr, f"{W}b")

    # Place the table at bookmark Spec_Table
    doc_el = doc.element
    body = doc_el.find(f".//{W}body")
    starts = [
        el
        for el in doc_el.findall(f".//{W}bookmarkStart")
        if el.get(f"{W}name") == "Spec_Table"
    ]
    if starts and body is not None:
        bm_start = starts[0]
        bm_id = bm_start.get(f"{W}id") or "26"
        bm_ends = [
            el
            for el in doc_el.findall(f".//{W}bookmarkEnd")
            if el.get(f"{W}id") == bm_id
        ]
        for end_el in bm_ends:
            if end_el.getparent() is not None:
                end_el.getparent().remove(end_el)
        if bm_start.getparent() is not None:
            anchor = bm_start
            while anchor.getparent() is not None and anchor.getparent() is not body:
                anchor = anchor.getparent()
            if anchor.getparent() is body:
                body.insert(list(body).index(anchor), table._tbl)
                body.remove(anchor)

        # Keep Spec_Table bookmark after <w:pPr> in cell(0, 0) so the bookmark remains valid
        first_p = table.rows[0].cells[0].paragraphs[0]._p
        pPr = first_p.find(f"{W}pPr")
        insert_idx = (list(first_p).index(pPr) + 1) if pPr is not None else 0
        new_start = OxmlElement("w:bookmarkStart")
        new_start.set(qn("w:id"), str(bm_id))
        new_start.set(qn("w:name"), "Spec_Table")
        new_end = OxmlElement("w:bookmarkEnd")
        new_end.set(qn("w:id"), str(bm_id))
        first_p.insert(insert_idx, new_end)
        first_p.insert(insert_idx, new_start)

    return table


def find_default_template(base_dir: Path = ROOT) -> Path:
    """Return the project's single supported default template."""
    template = base_dir / "ТКП_шаблон.dotx"
    if template.exists():
        return template
    raise FileNotFoundError(f"Word template ТКП_шаблон.dotx not found in {base_dir}")


def find_default_workbook(base_dir: Path = ROOT) -> Path:
    for name in ("Gen_TKP.xlsx", "Gen_TKP.xlsm"):
        candidate = base_dir / name
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Excel generator workbook not found in {base_dir}")


def populate_demo_workbook(wb: Workbook) -> None:
    """Populate a workbook in-memory with representative test data for demonstration."""
    info = _find_sheet(wb, "Наценки-Скидки")
    if info is not None:
        info["C2"] = "№ 104/1 от 16.10.2026"
        info["C3"] = "ООО «Газпром трансгаз Томск»"
        info["C4"] = "АСУ ТП компрессорного цеха КС-4"
        info["C5"] = (
            "Поставка оборудования и выполнение работ по созданию АСУ ТП "
            "на базе программно-технического комплекса АРКС400"
        )

    sw = _find_sheet(wb, "SW(w.o.VAT)", "ПО")
    if sw is not None:
        # Row 3: NVT software with broken #REF! in C3 -> falls back to D3 article
        sw["D3"] = "АРКС400-СДА"
        sw["F3"] = "шт."
        sw["G3"] = 1
        # Rows 14-15: InSAT software
        sw["C14"] = "Исполнительная система для программирования контроллеров. 1000 точек"
        sw["D14"] = "MasterPLC.N1000"
        sw["F14"] = "шт."
        sw["G14"] = 2
        sw["C15"] = "Лицензия на клиент визуализации. 1000 точек"
        sw["D15"] = "Визуализация.N1000"
        sw["F15"] = "шт."
        sw["G15"] = 2
        # Row 25: 3rd-party software
        sw["C25"] = "Операционная система Astra Linux Special Edition"
        sw["D25"] = "ALSE-1.7"
        sw["F25"] = "шт."
        sw["G25"] = 2

    cab1 = _find_sheet(wb, "1")
    if cab1 is not None:
        # Block 1 (НВТ -> Шкаф 1 ПТК): Row 3 tests #REF! fallback in C3
        cab1["D3"] = "АРКС400-ЦП-01"
        cab1["F3"] = "шт."
        cab1["G3"] = 2
        cab1["C4"] = "Модуль дискретного ввода 32 канала =24В"
        cab1["D4"] = "АРКС400-ДВ-32"
        cab1["F4"] = "шт."
        cab1["G4"] = 4
        cab1["C5"] = "Модуль аналогового ввода 16 каналов 4..20 мА"
        cab1["D5"] = "АРКС400-АВ-16"
        cab1["F5"] = "шт."
        cab1["G5"] = 3
        # Block 2 (Эксара -> Шкаф 1 телекоммуникационное оборудование)
        cab1["C14"] = "Коммутатор промышленный управляемый L2, 8x100/1000Base-T"
        cab1["D14"] = "EX-SW-8G"
        cab1["F14"] = "шт."
        cab1["G14"] = 2
        # Block 3 (3-х поставщиков -> Шкаф 1 шкафное оборудование)
        cab1["C20"] = "Шкаф напольный 2000х800х600 мм, IP54"
        cab1["D20"] = "TS-8806"
        cab1["F20"] = "шт."
        cab1["G20"] = 1
        cab1["C21"] = "Блок питания промышленный 24В DC, 20А"
        cab1["D21"] = "QUINT4-PS/1AC/24DC/20"
        cab1["F21"] = "шт."
        cab1["G21"] = 2

    cab2 = _find_sheet(wb, "2")
    if cab2 is not None:
        # Block 1 (НВТ -> Шкаф 2 ПТК)
        cab2["C3"] = "Модуль дискретного вывода 16 каналов =24В"
        cab2["D3"] = "АРКС400-ДВыв-16"
        cab2["F3"] = "шт."
        cab2["G3"] = 2
        # Block 2 (Эксара) deliberately left at 0 to verify empty block omission
        # Block 3 (3-х поставщиков -> Шкаф 2 шкафное оборудование)
        cab2["C20"] = "Шкаф навесной 800х600х300 мм, IP54"
        cab2["D20"] = "AE-1076"
        cab2["F20"] = "шт."
        cab2["G20"] = 1
        cab2["C21"] = "Клемма проходная винтовая 2,5 мм²"
        cab2["D21"] = "UT-2.5"
        cab2["F21"] = "шт."
        cab2["G21"] = 50

    sp = _find_sheet(wb, "Spare parts", "ЗИП")
    if sp is not None:
        # Row 3: NVT spare part with #REF! in C3 -> falls back to D3
        sp["D3"] = "АРКС400-ЦП-01"
        sp["F3"] = "шт."
        sp["G3"] = 1
        sp["C19"] = "Коммутатор промышленный управляемый L2, 8x100/1000Base-T"
        sp["D19"] = "EX-SW-8G"
        sp["F19"] = "шт."
        sp["G19"] = 1
        sp["C25"] = "Блок питания промышленный 24В DC, 20А"
        sp["D25"] = "QUINT4-PS/1AC/24DC/20"
        sp["F25"] = "шт."
        sp["G25"] = 1


def generate_tkp(
    excel_source: str | Path | Workbook | None = None,
    template_path: str | Path | None = None,
    output_path: str | Path | None = None,
    *,
    use_demo_if_empty: bool = False,
) -> TKPData:
    """Assemble the commercial proposal Word document from Excel data and template."""
    if excel_source is None:
        excel_source = find_default_workbook(ROOT)
    if template_path is None:
        template_path = find_default_template(ROOT)
    if output_path is None:
        output_path = ROOT / "Коммерческое предложение_результат.docx"

    data = read_tkp_data(excel_source)
    if data.total_items == 0 and use_demo_if_empty:
        if isinstance(excel_source, Workbook):
            wb = excel_source
        else:
            with warnings.catch_warnings():
                warnings.filterwarnings(
                    "ignore", message="Data Validation extension is not supported"
                )
                wb = openpyxl.load_workbook(excel_source, data_only=False)
        populate_demo_workbook(wb)
        update_preview_sheet(wb, read_tkp_data(wb).sections)
        data = read_tkp_data(wb)

    doc = open_word_template(template_path)
    doc_el = doc.element
    set_bookmark_text(doc_el, "Num", data.metadata.number)
    set_bookmark_text(doc_el, "Customer", data.metadata.customer)
    set_bookmark_text(doc_el, "Object", data.metadata.object_name)
    if data.metadata.subject:
        set_bookmark_text(doc_el, "Subject", data.metadata.subject)

    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            if hdr is not None:
                set_bookmark_text(hdr._element, "HeaderNum", data.metadata.number)

    populate_spec_table(doc, data.sections)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return data


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Сборка Коммерческого предложения (Приложение 3 и титульные реквизиты) из Gen_TKP."
    )
    parser.add_argument(
        "-e",
        "--excel",
        type=Path,
        default=None,
        help="Путь к файлу Gen_TKP.xlsx / Gen_TKP.xlsm",
    )
    parser.add_argument(
        "-t",
        "--template",
        type=Path,
        default=None,
        help="Путь к шаблону Word (.dotx / .docx)",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=ROOT / "Коммерческое предложение_результат.docx",
        help="Путь для сохранения готового документа Word",
    )
    args = parser.parse_args()

    data = generate_tkp(
        excel_source=args.excel,
        template_path=args.template,
        output_path=args.output,
        use_demo_if_empty=True,
    )
    print(
        f"Документ сохранен: {args.output} "
        f"(разделов: {len(data.sections)}, позиций: {data.total_items})"
    )


if __name__ == "__main__":
    main()
